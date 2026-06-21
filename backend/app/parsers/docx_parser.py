"""
Docx parser — parse .docx files using python-docx.
"""
import logging

from app.parsers.helpers import (
    ParseResult,
    _build_section_hierarchy,
    _extract_abbreviations_from_text,
    _extract_terms_from_text,
    _find_term_abbreviation_sections,
    _table_to_html,
)

logger = logging.getLogger(__name__)


class DocxParser:
    """Parse .docx files using python-docx."""

    def parse(self, file_path: str) -> ParseResult:
        """Parse a .docx file and return its structure."""
        try:
            from docx import Document as DocxDocument
        except ImportError:
            raise ImportError("python-docx is required to parse .docx files")

        doc = DocxDocument(file_path)
        result = ParseResult()
        section_counter = 0
        table_counter = 0

        # We need to interleave paragraphs and tables as they appear in the document.
        # python-docx's document.paragraphs and document.tables are separate lists,
        # so we iterate the XML body to preserve order.
        body = doc.element.body
        from docx.text.paragraph import Paragraph
        from docx.table import Table as DocxTable

        for child in body:
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag

            if tag == "p":
                paragraph = Paragraph(child, doc)
                self._process_paragraph(paragraph, result, section_counter)
                if result.sections:
                    section_counter = max(s["order_num"] for s in result.sections)

            elif tag == "tbl":
                try:
                    table = DocxTable(child, doc)
                    self._process_table(table, result, table_counter, section_counter)
                    table_counter += 1
                    if result.tables:
                        table_counter = max(t["order_num"] for t in result.tables) + 1
                except Exception as e:
                    logger.warning(f"Failed to parse table: {e}")

        # Build hierarchy
        result.sections = _build_section_hierarchy(result.sections)

        # Extract terms and abbreviations
        self._extract_terms_and_abbreviations(result)

        return result

    def _process_paragraph(self, paragraph, result: ParseResult, current_section: int) -> None:
        """Process a single paragraph."""
        style_name = paragraph.style.name if paragraph.style else "Normal"
        text = paragraph.text.strip()

        if not text:
            return

        if style_name and style_name.lower().startswith("heading"):
            level = 1
            parts = style_name.split()
            if len(parts) > 1:
                try:
                    level = int(parts[-1])
                except ValueError:
                    level = 1
            level = min(level, 6)

            result.sections.append({
                "title": text,
                "level": level,
                "content": "",
                "parent_id": None,
                "order_num": len(result.sections) + 1,
            })
        else:
            if result.sections:
                result.sections[-1]["content"] += text + "\n"

    def _process_table(self, table, result: ParseResult, table_counter: int, section_counter: int) -> None:
        """Process a single table."""
        headers = []
        rows_data = []

        for i, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells]
            if i == 0 and table_counter == 0:
                headers = cells
            else:
                rows_data.append(cells)

        if len(table.rows) <= 1 and headers:
            rows_data = [headers]
            headers = []

        section_id = section_counter if result.sections else None

        result.tables.append({
            "caption": None,
            "html_content": _table_to_html(headers, rows_data, None),
            "rows_count": len(rows_data) + (1 if headers else 0),
            "cols_count": max(
                (len(headers) if headers else 0),
                max((len(r) for r in rows_data), default=0),
            ),
            "order_num": len(result.tables) + 1,
            "section_id": section_id,
        })

        # Add table reference to current section
        if result.sections:
            result.sections[-1]["content"] += "\n[Таблица]\n"

    def _extract_terms_and_abbreviations(self, result: ParseResult) -> None:
        """Extract terms and abbreviations from parsed sections."""
        all_text = "\n".join(s.get("content", "") for s in result.sections)
        term_section_ids, abbr_section_ids = _find_term_abbreviation_sections(result.sections)

        # Extract terms from dedicated sections
        for section in result.sections:
            if section["order_num"] in term_section_ids:
                result.terms.extend(
                    _extract_terms_from_text(section.get("content", ""))
                )

        # Fallback: search entire document
        if not result.terms:
            result.terms.extend(_extract_terms_from_text(all_text))

        # Extract abbreviations
        for section in result.sections:
            if section["order_num"] in abbr_section_ids:
                result.abbreviations.extend(
                    _extract_abbreviations_from_text(section.get("content", ""))
                )

        if not result.abbreviations:
            result.abbreviations.extend(_extract_abbreviations_from_text(all_text))
