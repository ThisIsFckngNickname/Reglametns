"""
Text extractor — extracts text from uploaded draft files (docx/pdf/xlsx).
Preserves structure: headings, tables, lists.
"""

import logging
import os
import tempfile
from typing import Optional

from fastapi import UploadFile as FastAPIUploadFile

logger = logging.getLogger(__name__)

# Heading style names in Russian and English locale
_HEADING_STYLES = {
    "heading 1", "heading 2", "heading 3", "heading 4", "heading 5",
    "заголовок 1", "заголовок 2", "заголовок 3", "заголовок 4", "заголовок 5",
    "1", "2", "3", "4", "5",
}


class TextExtractor:
    """Extracts text content from uploaded files, preserving structure."""

    async def extract_draft_text(
        self,
        files: Optional[list[FastAPIUploadFile]],
    ) -> str:
        """Extract text from uploaded draft files (docx/pdf/xlsx)."""
        if not files:
            return ""

        texts = []
        for file in files:
            content = await file.read()
            filename = file.filename or "draft"
            ext = os.path.splitext(filename)[1].lower()

            if ext == ".docx":
                text = self._extract_docx(content)
            elif ext == ".pdf":
                text = self._extract_pdf(content)
            elif ext == ".xlsx":
                text = self._extract_xlsx(content)
            else:
                text = content.decode("utf-8", errors="ignore")

            if text:
                header = f"\n--- Файл: {filename} ---\n"
                texts.append(header + text)

        return "\n".join(texts)

    # ── .docx ──────────────────────────────────────────────────────────

    def _extract_docx(self, content: bytes) -> str:
        """Extract text from .docx preserving headings, tables, and lists."""
        try:
            from docx import Document as DocxDocument
        except ImportError:
            logger.warning("python-docx not installed")
            return ""

        with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
            tmp.write(content)
            tmp_path = tmp.name

        try:
            doc = DocxDocument(tmp_path)
            return self._render_docx(doc)
        except Exception as e:
            logger.warning(f"Failed to extract docx text: {e}")
            return ""
        finally:
            os.unlink(tmp_path)

    def _render_docx(self, doc) -> str:
        """Render .docx document to structured text.

        Walks body elements in document order to interleave
        paragraphs and tables correctly.
        """
        parts = []

        for element in doc.element.body:
            tag = element.tag.split("}")[-1]  # strip namespace

            if tag == "p":
                # Find matching paragraph in doc.paragraphs
                para = self._find_paragraph_by_el(doc, element)
                if para and para.text.strip():
                    rendered = self._render_paragraph(para)
                    if rendered:
                        parts.append(rendered)

            elif tag == "tbl":
                # Find matching table in doc.tables
                table = self._find_table_by_el(doc, element)
                if table:
                    rendered = self._render_table(table)
                    if rendered:
                        parts.append(rendered)

        return "\n".join(parts)

    def _find_paragraph_by_el(self, doc, el) -> Optional[object]:
        """Match lxml paragraph element back to python-docx Paragraph."""
        from docx.text.paragraph import Paragraph
        for p in doc.paragraphs:
            if p._element is el:
                return p
        return None

    def _find_table_by_el(self, doc, el) -> Optional[object]:
        """Match lxml table element back to python-docx Table."""
        for t in doc.tables:
            if t._element is el:
                return t
        return None

    def _render_paragraph(self, para) -> str:
        """Render a single paragraph, detecting headings."""
        style_name = (para.style.name or "").lower().strip()
        text = para.text.strip()
        if not text:
            return ""

        if style_name in _HEADING_STYLES:
            level = 0
            # Try to extract heading level from style name
            for ch in style_name:
                if ch.isdigit():
                    level = int(ch)
                    break
            prefix = "#" * max(1, level)
            return f"{prefix} {text}"

        # Detect bullet / numbered lists
        if para.style and "list" in style_name:
            return f"- {text}"
        if para.text.strip().startswith("- ") or para.text.strip().startswith("•"):
            return f"- {text[2:]}"

        return text

    def _render_table(self, table) -> str:
        """Render a .docx table as structured text."""
        rows = []
        for row_idx, row in enumerate(table.rows):
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if not rows:
            return ""

        lines = ["[ТАБЛИЦА]"]
        lines.append(rows[0])
        if len(rows) > 1:
            # Separator
            col_count = len(table.columns)
            lines.append(" | ".join(["---"] * col_count))
            lines.extend(rows[1:])
        lines.append("[/ТАБЛИЦА]")
        return "\n".join(lines)

    # ── .pdf ───────────────────────────────────────────────────────────

    def _extract_pdf(self, content: bytes) -> str:
        """Extract text from PDF, preserving some structure."""
        try:
            import fitz  # PyMuPDF
        except ImportError:
            logger.warning("PyMuPDF not installed")
            return ""

        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
            tmp.write(content)
            tmp_path = tmp.name

        try:
            doc = fitz.open(tmp_path)
            parts = []
            for page_num, page in enumerate(doc):
                blocks = page.get_text("dict", sort=True)["blocks"]
                page_parts = []

                for block in blocks:
                    if block["type"] == 0:  # text block
                        for line in block["lines"]:
                            text = "".join(
                                span["text"] for span in line["spans"]
                            ).strip()
                            if not text:
                                continue

                            # Detect headings by font size
                            font_sizes = [
                                span["size"]
                                for span in line["spans"]
                                if span["text"].strip()
                            ]
                            avg_size = (
                                sum(font_sizes) / len(font_sizes)
                                if font_sizes
                                else 0
                            )

                            if avg_size > 14:
                                page_parts.append(f"## {text}")
                            else:
                                page_parts.append(text)

                    elif block["type"] == 1:  # image block
                        pass  # skip images

                if page_parts:
                    parts.append(f"\n--- Страница {page_num + 1} ---")
                    parts.extend(page_parts)

            doc.close()
            return "\n".join(parts)

        except Exception as e:
            logger.warning(f"Failed to extract PDF text: {e}")
            return ""
        finally:
            os.unlink(tmp_path)

    # ── .xlsx ──────────────────────────────────────────────────────────

    def _extract_xlsx(self, content: bytes) -> str:
        """Extract text from .xlsx — reads all sheets."""
        try:
            import openpyxl
        except ImportError:
            logger.warning("openpyxl not installed, trying pandas")
            return self._extract_xlsx_pandas(content)

        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
            tmp.write(content)
            tmp_path = tmp.name

        try:
            wb = openpyxl.load_workbook(tmp_path, read_only=True, data_only=True)
            parts = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = []
                for row in ws.iter_rows(values_only=True):
                    vals = [str(c) if c is not None else "" for c in row]
                    if any(v.strip() for v in vals):
                        rows.append(" | ".join(vals))

                if rows:
                    parts.append(
                        f"\n--- Лист: {sheet_name} ---\n"
                        + "\n".join(rows)
                    )
            wb.close()
            return "\n".join(parts)

        except Exception as e:
            logger.warning(f"Failed to extract xlsx text via openpyxl: {e}")
            return ""
        finally:
            os.unlink(tmp_path)

    def _extract_xlsx_pandas(self, content: bytes) -> str:
        """Fallback: extract .xlsx via pandas."""
        try:
            import pandas as pd
        except ImportError:
            logger.warning("pandas not available for xlsx extraction")
            return ""

        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
            tmp.write(content)
            tmp_path = tmp.name

        try:
            dfs = pd.read_excel(tmp_path, sheet_name=None, dtype=str)
            parts = []
            for sheet_name, df in dfs.items():
                df = df.dropna(how="all")
                if df.empty:
                    continue
                rows = []
                for _, row in df.iterrows():
                    vals = [str(v) if pd.notna(v) else "" for v in row]
                    if any(v.strip() for v in vals):
                        rows.append(" | ".join(vals))
                if rows:
                    parts.append(
                        f"\n--- Лист: {sheet_name} ---\n"
                        + "\n".join(rows)
                    )
            return "\n".join(parts)

        except Exception as e:
            logger.warning(f"Failed to extract xlsx via pandas: {e}")
            return ""
        finally:
            os.unlink(tmp_path)


# Singleton
text_extractor = TextExtractor()
