"""
Document parser service.

Supports:
  - Word (.docx) via python-docx
  - PDF via PyMuPDF (fitz) + pdfplumber

Extracts:
  - Sections (headings, hierarchy)
  - Tables (HTML representation)
  - Terms and definitions
  - Abbreviations
"""

import logging
import os
import re
from dataclasses import dataclass, field
from typing import List, Optional

logger = logging.getLogger(__name__)


@dataclass
class ParseResult:
    """Result of parsing a document."""
    sections: List[dict] = field(default_factory=list)
    tables: List[dict] = field(default_factory=list)
    terms: List[dict] = field(default_factory=list)
    abbreviations: List[dict] = field(default_factory=list)
    lists: List[dict] = field(default_factory=list)  # NEW: detected lists


# ─── Term & Abbreviation Extraction ────────────────────────────────────

TERM_SECTION_KEYWORDS = [
    "термин", "определение", "глоссарий", "понятий",
    "термины и определения", "основные понятия",
]

ABBREVIATION_SECTION_KEYWORDS = [
    "сокращение", "условное обозначение", "аббревиатур",
    "список сокращений", "принятые сокращения",
]


def _is_term_section(title: str) -> bool:
    title_lower = title.lower().strip()
    return any(kw in title_lower for kw in TERM_SECTION_KEYWORDS)


def _is_abbreviation_section(title: str) -> bool:
    title_lower = title.lower().strip()
    return any(kw in title_lower for kw in ABBREVIATION_SECTION_KEYWORDS)


def _extract_terms_from_text(text: str) -> List[dict]:
    """Extract term-definition pairs from text.

    Heuristics:
      1. "Термин — определение" (em dash)
      2. "Термин – определение" (en dash)
      3. "Термин: определение" (colon)
      4. "Термин\tопределение" (tab)
    """
    terms = []
    lines = text.split("\n")

    for line in lines:
        line = line.strip()
        if not line:
            continue

        term = None
        definition = None

        for sep in [" — ", " – ", " - ", ": ", "\t"]:
            if sep in line:
                parts = line.split(sep, 1)
                candidate_term = parts[0].strip()
                candidate_def = parts[1].strip()

                if (
                    candidate_term
                    and candidate_def
                    and len(candidate_term) < 100
                    and not re.match(r"^\d", candidate_term)
                ):
                    term = candidate_term
                    definition = candidate_def
                    break

        if term and definition:
            term = term.rstrip(".,;:")
            definition = re.sub(r"^[\s\-–—:.]+\s*", "", definition)

            if term and definition:
                terms.append({"term": term, "definition": definition})

    return terms


def _extract_abbreviations_from_text(text: str) -> List[dict]:
    """Extract abbreviation-full_form pairs from text."""
    abbreviations = []
    lines = text.split("\n")

    for line in lines:
        line = line.strip()
        if not line:
            continue

        abbr = None
        full_form = None

        # Pattern 1: "ABBR — full form"
        for sep in [" — ", " – ", " - "]:
            if sep in line:
                parts = line.split(sep, 1)
                candidate_abbr = parts[0].strip()
                candidate_full = parts[1].strip()

                if (
                    candidate_abbr
                    and candidate_full
                    and len(candidate_abbr) < 30
                    and (candidate_abbr.isupper() or re.match(r"^[А-ЯA-Z]{2,}", candidate_abbr))
                ):
                    abbr = candidate_abbr
                    full_form = candidate_full
                    break

        # Pattern 2: "ABBR (full form)"
        if not abbr:
            match = re.match(r"^([А-ЯA-Z]{2,})\s*[\(（](.+)[\)）]", line)
            if match:
                abbr = match.group(1).strip()
                full_form = match.group(2).strip()

        if abbr and full_form:
            abbr = abbr.rstrip(".,;:")
            full_form = full_form.rstrip(".,;:")
            abbreviations.append({"abbreviation": abbr, "full_form": full_form})

    return abbreviations


def _find_term_abbreviation_sections(sections: List[dict]) -> tuple:
    """Find sections that contain terms or abbreviations."""
    term_section_ids = set()
    abbreviation_section_ids = set()

    for section in sections:
        title = section.get("title", "")
        if _is_term_section(title):
            term_section_ids.add(section.get("order_num"))
        if _is_abbreviation_section(title):
            abbreviation_section_ids.add(section.get("order_num"))

    return term_section_ids, abbreviation_section_ids


# ─── Section Hierarchy Builder ─────────────────────────────────────────

def _build_section_hierarchy(sections: List[dict]) -> List[dict]:
    """Build parent-child relationships based on heading levels."""
    if not sections:
        return []

    sections = sorted(sections, key=lambda s: s["order_num"])
    result = []
    parent_stack: list = []

    for section in sections:
        level = section.get("level", 1)
        order_num = section.get("order_num", 0)

        while parent_stack and parent_stack[-1][0] >= level:
            parent_stack.pop()

        if parent_stack:
            section["parent_id"] = parent_stack[-1][1]
        else:
            section["parent_id"] = None

        parent_stack.append((level, order_num))
        result.append(section)

    return result


# ─── HTML Table Builder ────────────────────────────────────────────────

def _table_to_html(headers: List[str], rows: List[List[str]], caption: Optional[str] = None) -> str:
    """Convert table data to simple HTML."""
    html_parts = []
    if caption:
        html_parts.append(f"<caption>{caption}</caption>")

    html_parts.append("<table border='1' cellpadding='4' cellspacing='0'>")

    if headers:
        html_parts.append("<thead><tr>")
        for h in headers:
            html_parts.append(f"<th>{h}</th>")
        html_parts.append("</tr></thead>")

    if rows:
        html_parts.append("<tbody>")
        for row in rows:
            html_parts.append("<tr>")
            for cell in row:
                html_parts.append(f"<td>{cell}</td>")
            html_parts.append("</tr>")
        html_parts.append("</tbody>")

    html_parts.append("</table>")
    return "\n".join(html_parts)


# ─── DOCX Parser ────────────────────────────────────────────────────────

class DocxParser:
    """Parse .docx files using python-docx."""

    def parse(self, file_path: str) -> ParseResult:
        """Parse a .docx file and return its structure."""
        try:
            from docx import Document as DocxDocument
        except ImportError:
            raise ImportError("python-docx is required to parse .docx files")

        doc = DocxDocument(file_path)
        result = ParseResult()
        section_counter = 0
        table_counter = 0

        # We need to interleave paragraphs and tables as they appear in the document.
        # python-docx's document.paragraphs and document.tables are separate lists,
        # so we iterate the XML body to preserve order.
        body = doc.element.body
        from docx.text.paragraph import Paragraph
        from docx.table import Table as DocxTable

        for child in body:
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag

            if tag == "p":
                paragraph = Paragraph(child, doc)
                self._process_paragraph(paragraph, result, section_counter)
                if result.sections:
                    section_counter = max(s["order_num"] for s in result.sections)

            elif tag == "tbl":
                try:
                    table = DocxTable(child, doc)
                    self._process_table(table, result, table_counter, section_counter)
                    table_counter += 1
                    if result.tables:
                        table_counter = max(t["order_num"] for t in result.tables) + 1
                except Exception as e:
                    logger.warning(f"Failed to parse table: {e}")

        # Build hierarchy
        result.sections = _build_section_hierarchy(result.sections)

        # Extract terms and abbreviations
        self._extract_terms_and_abbreviations(result)

        return result

    def _process_paragraph(self, paragraph, result: ParseResult, current_section: int) -> None:
        """Process a single paragraph."""
        style_name = paragraph.style.name if paragraph.style else "Normal"
        text = paragraph.text.strip()

        if not text:
            return

        if style_name and style_name.lower().startswith("heading"):
            level = 1
            parts = style_name.split()
            if len(parts) > 1:
                try:
                    level = int(parts[-1])
                except ValueError:
                    level = 1
            level = min(level, 6)

            result.sections.append({
                "title": text,
                "level": level,
                "content": "",
                "parent_id": None,
                "order_num": len(result.sections) + 1,
            })
        else:
            if result.sections:
                result.sections[-1]["content"] += text + "\n"

    def _process_table(self, table, result: ParseResult, table_counter: int, section_counter: int) -> None:
        """Process a single table."""
        headers = []
        rows_data = []

        for i, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells]
            if i == 0 and table_counter == 0:
                headers = cells
            else:
                rows_data.append(cells)

        if len(table.rows) <= 1 and headers:
            rows_data = [headers]
            headers = []

        section_id = section_counter if result.sections else None

        result.tables.append({
            "caption": None,
            "html_content": _table_to_html(headers, rows_data, None),
            "rows_count": len(rows_data) + (1 if headers else 0),
            "cols_count": max(
                (len(headers) if headers else 0),
                max((len(r) for r in rows_data), default=0),
            ),
            "order_num": len(result.tables) + 1,
            "section_id": section_id,
        })

        # Add table reference to current section
        if result.sections:
            result.sections[-1]["content"] += "\n[Таблица]\n"

    def _extract_terms_and_abbreviations(self, result: ParseResult) -> None:
        """Extract terms and abbreviations from parsed sections."""
        all_text = "\n".join(s.get("content", "") for s in result.sections)
        term_section_ids, abbr_section_ids = _find_term_abbreviation_sections(result.sections)

        # Extract terms from dedicated sections
        for section in result.sections:
            if section["order_num"] in term_section_ids:
                result.terms.extend(
                    _extract_terms_from_text(section.get("content", ""))
                )

        # Fallback: search entire document
        if not result.terms:
            result.terms.extend(_extract_terms_from_text(all_text))

        # Extract abbreviations
        for section in result.sections:
            if section["order_num"] in abbr_section_ids:
                result.abbreviations.extend(
                    _extract_abbreviations_from_text(section.get("content", ""))
                )

        if not result.abbreviations:
            result.abbreviations.extend(_extract_abbreviations_from_text(all_text))


# ─── PDF Parser ─────────────────────────────────────────────────────────

class PdfParser:
    """Parse PDF files using PyMuPDF (fitz) + pdfplumber."""

    def parse(self, file_path: str) -> ParseResult:
        """Parse a PDF file and return its structure."""
        result = ParseResult()
        all_text = ""

        # Step 1: Extract text with PyMuPDF
        try:
            import fitz
        except ImportError:
            raise ImportError("PyMuPDF (fitz) is required to parse PDF files")

        try:
            doc = fitz.open(file_path)
        except Exception as e:
            logger.warning(f"Failed to open PDF with PyMuPDF: {e}")
            return result

        for page_num, page in enumerate(doc):
            blocks = page.get_text("dict").get("blocks", [])
            page_text_lines = []

            for block in blocks:
                if block.get("type") != 0:  # skip images
                    continue

                for line in block.get("lines", []):
                    line_text = ""
                    font_size = None
                    is_bold = False

                    for span in line.get("spans", []):
                        line_text += span.get("text", "")
                        if font_size is None:
                            font_size = span.get("size", 12)
                        font_name = span.get("font", "").lower()
                        if "bold" in font_name:
                            is_bold = True

                    line_text = line_text.strip()
                    if not line_text:
                        continue

                    page_text_lines.append(line_text)

                    # Detect heading: large font (>14) or bold+large
                    is_heading = (
                        (font_size and font_size > 14)
                        or (is_bold and font_size and font_size > 12)
                        or (line_text.isupper() and len(line_text) > 3 and len(line_text) < 200)
                    )

                    if is_heading and len(line_text) < 200:
                        level = 1
                        if font_size:
                            if font_size <= 14:
                                level = 2
                            elif font_size <= 18:
                                level = 1
                            else:
                                level = 1

                        result.sections.append({
                            "title": line_text,
                            "level": level,
                            "content": "",
                            "parent_id": None,
                            "order_num": len(result.sections) + 1,
                        })
                    elif result.sections:
                        result.sections[-1]["content"] += line_text + "\n"

            all_text += "\n".join(page_text_lines) + "\n"

        doc.close()

        # Step 2: Extract tables with pdfplumber
        try:
            self._extract_pdf_tables(file_path, result)
        except Exception as e:
            logger.warning(f"Failed to extract tables with pdfplumber: {e}")

        # Build hierarchy
        result.sections = _build_section_hierarchy(result.sections)

        # Extract terms and abbreviations
        term_section_ids, abbr_section_ids = _find_term_abbreviation_sections(result.sections)

        for section in result.sections:
            if section["order_num"] in term_section_ids:
                result.terms.extend(
                    _extract_terms_from_text(section.get("content", ""))
                )

        if not result.terms:
            result.terms.extend(_extract_terms_from_text(all_text))

        for section in result.sections:
            if section["order_num"] in abbr_section_ids:
                result.abbreviations.extend(
                    _extract_abbreviations_from_text(section.get("content", ""))
                )

        if not result.abbreviations:
            result.abbreviations.extend(_extract_abbreviations_from_text(all_text))

        return result

    def _extract_pdf_tables(self, file_path: str, result: ParseResult) -> None:
        """Extract tables from PDF using pdfplumber."""
        try:
            import pdfplumber
        except ImportError:
            logger.warning("pdfplumber not available, skipping table extraction")
            return

        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                tables = page.extract_tables()
                for table_data in tables:
                    if not table_data:
                        continue

                    headers = []
                    rows = []
                    for i, row in enumerate(table_data):
                        cells = [str(c) if c else "" for c in row]
                        if i == 0:
                            headers = cells
                        else:
                            rows.append(cells)

                    if len(table_data) <= 1 and headers:
                        rows = [headers]
                        headers = []

                    section_id = result.sections[-1]["order_num"] if result.sections else None

                    result.tables.append({
                        "caption": None,
                        "html_content": _table_to_html(headers, rows, None),
                        "rows_count": len(rows) + (1 if headers else 0),
                        "cols_count": max(
                            (len(headers) if headers else 0),
                            max((len(r) for r in rows), default=0),
                        ),
                        "order_num": len(result.tables) + 1,
                        "section_id": section_id,
                    })


# ─── Document Parser Facade ─────────────────────────────────────────────

class DocumentParser:
    """Facade that picks the right parser based on file type."""

    def __init__(self, enhanced: bool = False):
        self._enhanced = enhanced

    def parse(self, file_path: str, file_type: str) -> ParseResult:
        if file_type == "pdf" and self._enhanced:
            try:
                from app.parsers.enhanced_pdf_parser import EnhancedPDFParser
                return EnhancedPDFParser().parse(file_path)
            except ImportError:
                logger.warning("EnhancedPDFParser not available, falling back to default PDF parser")

        if file_type == "docx":
            return DocxParser().parse(file_path)
        elif file_type == "pdf":
            return PdfParser().parse(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")


# Convenience instances
document_parser = DocumentParser(enhanced=False)
enhanced_document_parser = DocumentParser(enhanced=True)


# ─── Order Parser ────────────────────────────────────────────────────────

@dataclass
class OrderParseResult:
    """Result of parsing an order document."""
    order_number: Optional[str] = None
    order_date: Optional[str] = None
    affected_documents: List[str] = field(default_factory=list)


class OrderParser:
    """
    Parses an order document to extract:
    - order_number: номер приказа (ищет "ПРИКАЗ №..." или "№...")
    - order_date: дата (ищет "от dd.mm.yyyy" или "dd.mm.yyyy")
    - affected_documents: список упомянутых документов (по названиям)
    """

    # Patterns for order number
    ORDER_NUMBER_PATTERNS = [
        re.compile(r'(?:ПРИКАЗ|Приказ|приказ)\s*[№#]?\s*([\d\-/]+)', re.UNICODE),
        re.compile(r'[№#]\s*([\d\-/]+)', re.UNICODE),
    ]

    # Patterns for order date
    ORDER_DATE_PATTERNS = [
        re.compile(r'от\s+(\d{2}\.\d{2}\.\d{4})', re.UNICODE),
        re.compile(r'(\d{2}\.\d{2}\.\d{4})\s*г', re.UNICODE),
        re.compile(r'"(\d{2})"\s*(\w+)\s*(\d{4})', re.UNICODE),
    ]

    # Russian month names for date parsing
    MONTH_MAP = {
        "января": "01", "февраля": "02", "марта": "03",
        "апреля": "04", "мая": "05", "июня": "06",
        "июля": "07", "августа": "08", "сентября": "09",
        "октября": "10", "ноября": "11", "декабря": "12",
    }

    def parse_order_details(self, file_path: str, file_type: str) -> OrderParseResult:
        """
        Parse an order file and extract details.

        Returns OrderParseResult with extracted fields.
        """
        result = OrderParseResult()

        # Extract text from file
        if file_type == "docx":
            text = self._extract_text_from_docx(file_path)
        elif file_type == "pdf":
            text = self._extract_text_from_pdf(file_path)
        else:
            logger.warning(f"Unsupported file type for order parsing: {file_type}")
            return result

        if not text:
            return result

        # Extract order number
        result.order_number = self._extract_order_number(text)

        # Extract order date
        date_str = self._extract_order_date(text)
        if date_str:
            # Try to normalize to YYYY-MM-DD
            normalized = self._normalize_date(date_str)
            if normalized:
                result.order_date = normalized

        # Extract affected documents
        result.affected_documents = self._extract_affected_documents(text)

        return result

    def _extract_text_from_docx(self, file_path: str) -> str:
        """Extract text from a .docx file."""
        try:
            from docx import Document as DocxDocument
            doc = DocxDocument(file_path)
            return "\n".join(p.text for p in doc.paragraphs)
        except ImportError:
            logger.error("python-docx is required to parse .docx files")
            return ""
        except Exception as e:
            logger.warning(f"Failed to extract text from docx: {e}")
            return ""

    def _extract_text_from_pdf(self, file_path: str) -> str:
        """Extract text from a PDF file."""
        try:
            import fitz
            doc = fitz.open(file_path)
            text = ""
            for page in doc:
                text += page.get_text()
            doc.close()
            return text
        except ImportError:
            logger.error("PyMuPDF (fitz) is required to parse PDF files")
            return ""
        except Exception as e:
            logger.warning(f"Failed to extract text from pdf: {e}")
            return ""

    def _extract_order_number(self, text: str) -> Optional[str]:
        """Extract order number from text."""
        for pattern in self.ORDER_NUMBER_PATTERNS:
            match = pattern.search(text)
            if match:
                return match.group(1).strip()
        return None

    def _extract_order_date(self, text: str) -> Optional[str]:
        """Extract order date string from text."""
        # Try dd.mm.yyyy patterns first
        for pattern in self.ORDER_DATE_PATTERNS[:2]:
            match = pattern.search(text)
            if match:
                return match.group(1).strip()

        # Try Russian text date: "dd month yyyy"
        match = self.ORDER_DATE_PATTERNS[2].search(text)
        if match:
            day = match.group(1)
            month_text = match.group(2).lower().strip(".")
            year = match.group(3)
            month_num = self.MONTH_MAP.get(month_text)
            if month_num:
                return f"{day}.{month_num}.{year}"

        return None

    def _normalize_date(self, date_str: str) -> Optional[str]:
        """Normalize date string to YYYY-MM-DD format."""
        import re
        match = re.match(r'(\d{2})\.(\d{2})\.(\d{4})', date_str)
        if match:
            return f"{match.group(3)}-{match.group(2)}-{match.group(1)}"
        return None

    def _extract_affected_documents(self, text: str) -> List[str]:
        """Extract names of potentially affected documents from text."""
        # Simple heuristic: look for quoted document titles
        affected = []
        lines = text.split("\n")
        for line in lines:
            # Look for "внести изменения в..." or "отменить..."
            lower = line.lower().strip()
            if any(kw in lower for kw in [
                "внести изменения", "отменить", "признать утратившим",
                "изложить в новой", "дополнить",
            ]):
                # Extract quoted text
                import re
                quoted = re.findall(r'"([^"]+)"', line)
                affected.extend(quoted)

        return affected


# Convenience instance
order_parser = OrderParser()
